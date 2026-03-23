import os
from fastapi import FastAPI, UploadFile, File, Query
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv
from test.rag_chain import load_rag_chain, get_answer
from sarvam_transcribe import transcribe_chunk, transcribe_file, generate_meeting_notes

from fastapi.responses import StreamingResponse
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.units import inch
import io
load_dotenv()

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
    allow_credentials=False,
)

print("Loading RAG chain...")
chain_tuple = load_rag_chain()
print("Ready!")

# In-memory session store
meeting_transcripts = {}

class QuestionRequest(BaseModel):
    question: str

class EndMeetingRequest(BaseModel):
    session_id: str
    meeting_type: str = "internal"

# ── Document Q&A ──
@app.post("/ask")
async def ask_question(req: QuestionRequest):
    answer, sources, _ = get_answer(chain_tuple, req.question)
    return {"answer": answer, "sources": sources}

# ── Upload doc ──
@app.post("/upload-doc")
async def upload_document(file: UploadFile = File(...)):
    os.makedirs("data", exist_ok=True)
    path = f"data/{file.filename}"
    with open(path, "wb") as f:
        f.write(await file.read())
    os.system("python ingest.py")
    return {"message": f"{file.filename} uploaded and indexed successfully"}

# ── Live meeting chunk ──
@app.post("/meeting/transcribe-chunk")
async def transcribe_meeting_chunk(
    session_id: str = Query(...),
    language: str = Query("en-IN"),
    file: UploadFile = File(...)
):
    audio_bytes = await file.read()
    print(f"[CHUNK] Session: {session_id} | Size: {len(audio_bytes)} bytes | Lang: {language}")

    # Convert webm to wav bytes for Sarvam
    chunk_transcript = transcribe_chunk(audio_bytes, language)
    print(f"[CHUNK] Transcript: '{chunk_transcript}'")

    if session_id not in meeting_transcripts:
        meeting_transcripts[session_id] = []

    if chunk_transcript.strip():
        meeting_transcripts[session_id].append(chunk_transcript)

    total = len(meeting_transcripts.get(session_id, []))
    print(f"[CHUNK] Total chunks so far: {total}")

    return {
        "chunk_transcript": chunk_transcript,
        "full_transcript": " ".join(meeting_transcripts.get(session_id, []))
    }

# ── End meeting → generate notes ──
@app.post("/meeting/end")
async def end_meeting(req: EndMeetingRequest):
    session_id = req.session_id
    print(f"[END] Session: {session_id}")

    if session_id not in meeting_transcripts:
        return {"error": "No transcript found for this session."}

    full_transcript = " ".join(meeting_transcripts[session_id])
    print(f"[END] Full transcript length: {len(full_transcript)} chars")

    if not full_transcript.strip():
        return {"error": "No speech detected during the meeting."}

    notes = generate_meeting_notes(full_transcript, req.meeting_type)
    notes["full_transcript"] = full_transcript
    del meeting_transcripts[session_id]
    return notes

# ── Upload existing recording file ──
@app.post("/meeting/transcribe-file")
async def transcribe_meeting_file(
    language: str = Query("en-IN"),
    meeting_type: str = Query("internal"),
    file: UploadFile = File(...)
):
    os.makedirs("temp", exist_ok=True)
    path = f"temp/{file.filename}"
    with open(path, "wb") as f:
        f.write(await file.read())

    transcript = transcribe_file(path, language)
    notes = generate_meeting_notes(transcript, meeting_type)
    notes["full_transcript"] = transcript
    return notes

@app.post("/meeting/download/docx")
async def download_docx(data: dict):
    summary = data.get("summary", "")
    key_decisions = data.get("key_decisions", "")
    action_items = data.get("action_items", "")
    transcript = data.get("transcript", "")

    doc = Document()
    doc.add_heading("KnowBot — Meeting Summary", 0)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(summary if summary else "No summary available.")

    doc.add_heading("Key Decisions", level=1)
    doc.add_paragraph(key_decisions if key_decisions else "No key decisions recorded.")

    doc.add_heading("Action Items", level=1)
    doc.add_paragraph(action_items if action_items else "No action items recorded.")

    doc.add_heading("Full Transcript", level=1)
    doc.add_paragraph(transcript if transcript else "No transcript available.")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=meeting-summary.docx"}
    )


@app.post("/meeting/download/pdf")
async def download_pdf(data: dict):
    summary = data.get("summary", "")
    key_decisions = data.get("key_decisions", "")
    action_items = data.get("action_items", "")
    transcript = data.get("transcript", "")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                            rightMargin=inch, leftMargin=inch,
                            topMargin=inch, bottomMargin=inch)

    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("KnowBot — Meeting Summary", styles['Title']))
    story.append(Spacer(1, 0.2 * inch))

    story.append(Paragraph("Summary", styles['Heading1']))
    story.append(Paragraph(summary if summary else "No summary available.", styles['Normal']))
    story.append(Spacer(1, 0.2 * inch))

    story.append(Paragraph("Key Decisions", styles['Heading1']))
    story.append(Paragraph(key_decisions if key_decisions else "No key decisions recorded.", styles['Normal']))
    story.append(Spacer(1, 0.2 * inch))

    story.append(Paragraph("Action Items", styles['Heading1']))
    story.append(Paragraph(action_items if action_items else "No action items recorded.", styles['Normal']))
    story.append(Spacer(1, 0.2 * inch))

    story.append(Paragraph("Full Transcript", styles['Heading1']))
    story.append(Paragraph(transcript if transcript else "No transcript available.", styles['Normal']))

    doc.build(story)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=meeting-summary.pdf"}
    )